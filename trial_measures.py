"""Reaction time, movement time etc. per trial - the same method as the previous (MATLAB) system - and the
Excel output.

Port of loaddata.m + calculateMeasures.m + makeCSVsubject.m / makeCSVall.m:
  - each trial's data runs from the end of the previous trial (the home phase) to the end of this trial
  - the marker (wrist for the upper limb, ankle for the lower limb - 3D with two cameras, otherwise its
    image position) is resampled to 100 Hz, low-pass filtered (Butterworth, 2nd order, 10 Hz, zero-phase),
    differentiated, and the tangential velocity is the speed
  - first peak = the first speed peak (> 30% of the maximum after the light came on) after the light came on
  - movement onset = the last sample between the light coming on and the first peak with speed < 5% of the
    maximum (the light-on sample if there is none) -> reaction time (responsetime) = onset - light on
  - movement end = the first sample after the first peak with speed < 5% of the maximum (the last sample if
    there is none) -> movement time = end - onset
  - also: straight-line deviation, number of speed peaks (> 10% of the maximum), acceleration time (AT),
    deceleration time (DT), AT/DT, end-point distance to the target
  - summaries: mean, median, SD and IQR by part, by block and by sequence, without familiarization, only
    the trials with touse = 1 (MATLAB's std and iqr: sample SD, 'hazen' percentiles)
Flags (new):
  - Anticipation: the movement toward the target started before the light came on (from the home position)
  - Corrected: the first movement went to another switch, then the right one was hit in time
"""
import csv
import os
import warnings

import numpy as np
from scipy.signal import butter, filtfilt, find_peaks

RESAMPLE_HZ = 100
FILTER_B, FILTER_A = butter(2, 10 / (RESAMPLE_HZ / 2))
MARKER = {'upper': {'right': 16, 'left': 15}, 'lower': {'right': 28, 'left': 27}}  # wrist / ankle
MIN_VISIBILITY = 0.3
ANTICIPATION_MIN_S = 0.03        # movement onset at least this much before the light -> anticipation
ANTICIPATION_MAX_S = 0.5         # ... and at most this much (earlier = still the return to home, not a reach)
HOME_ZONE = 0.3                  # ... and it started within 30% of the home-target distance from home
CORRECTION_MIN_S = 0.15          # first movement stopped at least this long before the hit ...
EXCLUDED_FROM_SUMMARY = ['Familiarization']
SUMMARY_PARTS = ['pre_test', 'pre_test_contralateral', 'practice', 'post_test', 'post_test_contralateral',
                 're_test', 're_test_contralateral']
MEASURES = ['responsetime', 'movementtime', 'straighlinedeviation', 'numpeaks', 'AT', 'DT', 'AT_over_DT',
            'endpointdistance']


SPIKE_FRACTION = 0.03            # a point > 3% of the image size away from its neighbours' median is a glitch


def marker_track(pose, limb, side, image_size=None):
    """(frames, 2) marker position (NaN where MediaPipe isn't sure, or where it jumps for a single frame or two
    - a glitch, not a movement: compared with the median of the 5 frames around it)."""
    j = MARKER[limb][side]
    xy = pose[:, j, :2].astype(float).copy()
    xy[~(pose[:, j, 2] >= MIN_VISIBILITY)] = np.nan
    if image_size is not None and len(xy) >= 5:
        pad = np.pad(xy, ((2, 2), (0, 0)), mode='edge')
        with warnings.catch_warnings():   # all-NaN stretches (marker not seen) are expected
            warnings.simplefilter('ignore', RuntimeWarning)
            med = np.nanmedian(np.lib.stride_tricks.sliding_window_view(pad, 5, axis=0), axis=2)
        jump = np.linalg.norm(xy - med, axis=1) > SPIKE_FRACTION * max(image_size)
        xy[jump] = np.nan
    return xy


def _point_line_distance(pts, v1, v2):
    """point_to_line_distance.m: distance of each point to the line through v1 and v2 (2D or 3D)."""
    pad = lambda a: np.hstack([a, np.zeros((a.shape[0], 3 - a.shape[1]))]) if a.shape[1] == 2 else a
    p, a, b = pad(np.atleast_2d(pts)), pad(np.atleast_2d(v1)), pad(np.atleast_2d(v2))
    d = a - b
    n = np.linalg.norm(d)
    return np.linalg.norm(np.cross(d, p - b), axis=1) / n if n > 0 else np.linalg.norm(p - b, axis=1)


def trial_kinematics(t_s, pos, stim_s):
    """loaddata.m for one trial: 100 Hz time, position, filtered speed. None if too little data."""
    ok = ~np.isnan(pos[:, 0])
    if ok.sum() < 5 or ok.mean() < 0.5:
        return None
    time = np.arange(t_s[0], t_s[-1] + 1e-9, 1 / RESAMPLE_HZ)
    xyz = np.column_stack([np.interp(time, t_s[ok], pos[ok, c]) for c in range(pos.shape[1])])
    padlen = 3 * (max(len(FILTER_A), len(FILTER_B)) - 1)
    if len(time) <= padlen:
        return None
    filt = filtfilt(FILTER_B, FILTER_A, xyz, axis=0, padlen=padlen)
    vel = np.vstack([np.zeros((1, xyz.shape[1])), np.diff(filt, axis=0) * RESAMPLE_HZ])
    return dict(time=time, xyz=xyz, tangvel=np.linalg.norm(vel, axis=1), vel=vel, stim=stim_s)


def calculate_measures(k, target_pos, home_pos=None):
    """calculateMeasures.m for one trial (k = trial_kinematics output). Returns measures + plot samples."""
    nan = dict.fromkeys(MEASURES, np.nan)
    if k is None:
        return dict(nan, onset_s=np.nan, end_s=np.nan, anticipation=False, anticipation_ms=None)
    time, xyz, tv, stim = k['time'], k['xyz'], k['tangvel'], k['stim']
    s = int(np.searchsorted(time, stim))   # MATLAB: find(time >= stimulusonsettime, 1)
    if s >= len(time) - 1:
        return dict(nan, onset_s=np.nan, end_s=np.nan, anticipation=False, anticipation_ms=None)
    maxtv = tv[s:].max()
    peaks, _ = find_peaks(tv, height=0.3 * maxtv)
    after = peaks[time[peaks] >= stim]
    if maxtv <= 0 or len(after) == 0:
        return dict(nan, onset_s=np.nan, end_s=np.nan, anticipation=False, anticipation_ms=None)
    first = int(after[0])
    below = np.nonzero(tv[s:first + 1] < 0.05 * maxtv)[0]
    onset = s + int(below[-1]) if len(below) else s
    below = np.nonzero(tv[first:] < 0.05 * maxtv)[0]
    end = first + int(below[0]) if len(below) else len(time) - 1
    rng = slice(onset, end + 1)
    seg = xyz[rng]
    m = dict(responsetime=time[onset] - stim, movementtime=time[end] - time[onset],
             straighlinedeviation=float(np.mean(_point_line_distance(seg, seg[0], seg[-1]))),
             numpeaks=len(find_peaks(tv[rng], height=0.1 * tv.max())[0]),
             AT=time[first] - time[onset], DT=time[end] - time[first])
    m['AT_over_DT'] = m['AT'] / m['DT'] if m['DT'] > 0 else np.nan
    m['endpointdistance'] = float(np.linalg.norm(seg[-1] - target_pos)) if target_pos is not None else np.nan
    m.update(onset_s=time[onset], end_s=time[end], peak_s=time[first])

    # Anticipation: the same onset search, but not stopped at the light: did the movement that peaks first
    # after the light start before it, from the home position, toward the target?
    below = np.nonzero(tv[:first + 1] < 0.05 * maxtv)[0]
    free_onset = int(below[-1]) if len(below) else 0
    m['anticipation'], m['anticipation_ms'] = False, None
    lead = stim - time[free_onset]
    if ANTICIPATION_MIN_S < lead <= ANTICIPATION_MAX_S and home_pos is not None and target_pos is not None:
        u = np.asarray(target_pos) - home_pos
        dist = np.linalg.norm(u)
        if dist > 0:
            u = u / dist
            near_home = np.linalg.norm(xyz[free_onset] - home_pos) < HOME_ZONE * dist
            toward = np.dot(k['vel'][s], u) > 0
            if near_home and toward:
                m['anticipation'] = True
                m['anticipation_ms'] = round((stim - time[free_onset]) * 1000, 1)
    return m


def nearest_switch(p, switch_pos, switch_size):
    d = np.linalg.norm(switch_pos - p, axis=1) / switch_size
    return int(np.argmin(d)) + 1


def corrected_reach(k, m, hit_s, target, switch_pos, switch_size):
    """Which other switch the first movement went to, if the target was then hit after a correction."""
    if k is None or np.isnan(m.get('end_s', np.nan)) or hit_s is None:
        return None
    time, xyz = k['time'], k['xyz']
    first_end = nearest_switch(xyz[np.searchsorted(time, m['end_s'])], switch_pos, switch_size)
    at_hit = nearest_switch(xyz[min(len(time) - 1, np.searchsorted(time, hit_s))], switch_pos, switch_size)
    if hit_s - m['end_s'] >= CORRECTION_MIN_S and first_end != at_hit and first_end != target:
        return first_end
    return None


def _homes(trials, track):
    """Home position per limb/side: where the marker is when the lights come on (median over trials)."""
    homes = {}
    for key, tr in track.items():
        on = [r['on_frame'] for r in trials if r['on_frame'] and (r['Limb'], r['Side']) == key]
        pts = tr[on] if on else np.empty((0, tr.shape[1]))
        pts = pts[~np.isnan(pts[:, 0])]
        homes[key] = np.median(pts, axis=0) if len(pts) else None
    return homes


def measure_trials(trials, times_ms, track, switch_pos, switch_size, source='camera A', fallback=None):
    """Add measures and flags to the trial rows (in place). `track`: (frames, 2 or 3) marker positions per
    limb/side: {(limb, side): array}. switch_pos: (switches, dims) in the same coordinates.
    fallback = (track, switch_pos, switch_size, source) used for a trial without enough data in `track`
    (e.g. 3D needs both cameras to see the marker); distances are then left empty (other units)."""
    t_s = np.asarray(times_ms) / 1000
    homes = _homes(trials, track)
    fb_homes = _homes(trials, fallback[0]) if fallback else {}
    for r in trials:
        r.update(dict.fromkeys(MEASURES, np.nan), anticipation=False, anticipation_ms=None, corrected_from=None,
                 touse=r['Outcome'] == 'hit')
        if r['on_frame'] is None or r.get('off_frame') is None:
            r['touse'] = False
            continue
        key = (r['Limb'], r['Side'])
        a, b = r['window_start'], r['off_frame']
        tr, sp, ss, home, r['Data'] = track, switch_pos, switch_size, homes.get(key), source
        k = trial_kinematics(t_s[a:b + 1], tr[key][a:b + 1], t_s[r['on_frame']])
        m = calculate_measures(k, sp[r['Target'] - 1], home)
        if np.isnan(m['responsetime']) and fallback:   # too little / too patchy data: use the fallback
            tr, sp, ss, r['Data'] = fallback
            home = fb_homes.get(key)
            k = trial_kinematics(t_s[a:b + 1], tr[key][a:b + 1], t_s[r['on_frame']])
            m = calculate_measures(k, sp[r['Target'] - 1], home)
        if r['Data'] != source:
            m['straighlinedeviation'] = m['endpointdistance'] = np.nan
        r.update({x: m.get(x, np.nan) for x in MEASURES})
        r['anticipation'], r['anticipation_ms'] = m['anticipation'], m['anticipation_ms']
        r['onset_s'], r['end_s'] = m.get('onset_s'), m.get('end_s')
        if r['Outcome'] == 'hit':
            r['corrected_from'] = corrected_reach(k, m, t_s[r['off_frame']], r['Target'], sp, ss)
        if k is None:
            r['touse'] = False
    return trials


# -----------------------------
# Output files
# -----------------------------
TRIAL_COLUMNS = ['Part', 'Block', 'Sequence', 'Trial', 'Target', 'Limb', 'Side', 'Light_on_ms', 'Touch_ms',
                 'Hit_ms', 'Light_off_ms', 'Trial_end_ms', 'Outcome', 'Anticipation', 'Anticipation_ms',
                 'Corrected', 'Corrected_from']


def _cell(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ''
    if isinstance(v, (bool, np.bool_)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        return round(float(v), 4)
    if isinstance(v, np.integer):
        return int(v)
    return v


def write_trials_csv(path, trials):
    """The trials synced with the experimental design, with the flags (one row per trial)."""
    with open(path, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(TRIAL_COLUMNS)
        for r in trials:
            w.writerow([_cell(r.get(c)) for c in TRIAL_COLUMNS[:13]] +
                       [_cell(r.get('anticipation')), _cell(r.get('anticipation_ms')),
                        int(r.get('corrected_from') is not None), _cell(r.get('corrected_from'))])


def _iqr(x):
    return np.percentile(x, 75, method='hazen') - np.percentile(x, 25, method='hazen') if len(x) else np.nan


STATS = {'mean': np.mean, 'median': np.median, 'sd': lambda x: np.std(x, ddof=1) if len(x) > 1 else np.nan,
         'iqr': _iqr}


def summary_tables(rows_by_subject):
    """makeCSVall.m: 12 wide tables (mean/median/sd/iqr x by part/block/sequence), one row per subject.
    rows_by_subject: {subject: [measure rows with Part/Block/Sequence/touse + MEASURES]}."""
    tables = {f'{stat}_by{level}': [] for level in ('part', 'block', 'sequence') for stat in STATS}
    for subject, rows in rows_by_subject.items():
        use = [r for r in rows if r['touse'] and r['Part'] not in EXCLUDED_FROM_SUMMARY]
        for stat, fn in STATS.items():
            out = {lvl: {'subject': subject} for lvl in ('part', 'block', 'sequence')}
            for meas in MEASURES:
                for part in SUMMARY_PARTS:
                    pr = [r for r in use if r['Part'] == part]
                    if not pr:
                        continue
                    vals = lambda rr: np.array([r[meas] for r in rr if not np.isnan(r[meas])], float)
                    v = vals(pr)
                    out['part'][f'{stat}_{part}_{meas}'] = fn(v) if len(v) else np.nan
                    for b in sorted({r['Block'] for r in pr}):
                        br = [r for r in pr if r['Block'] == b]
                        v = vals(br)
                        out['block'][f'{stat}_{part}_block{b}_{meas}'] = fn(v) if len(v) else np.nan
                        for sq in sorted({r['Sequence'] for r in br}):
                            v = vals([r for r in br if r['Sequence'] == sq])
                            out['sequence'][f'{stat}_{part}_block{b}_sequence{sq}_{meas}'] = fn(v) if len(v) else np.nan
            for lvl in out:
                tables[f'{stat}_by{lvl}'].append(out[lvl])
    return tables


MEASURE_COLUMNS = (['Subject', 'Session', 'Part', 'Block', 'Sequence', 'Trial', 'Target', 'Outcome', 'Anticipation',
                    'Anticipation_ms', 'Corrected', 'Corrected_from', 'Data'] + MEASURES + ['touse'])


def measure_row(subject, session, r):
    return {'Subject': subject, 'Session': session, 'Part': r['Part'], 'Block': r['Block'],
            'Sequence': r['Sequence'], 'Trial': r['Trial'], 'Target': r['Target'], 'Outcome': r['Outcome'],
            'Anticipation': bool(r.get('anticipation')), 'Anticipation_ms': r.get('anticipation_ms'),
            'Corrected': r.get('corrected_from') is not None, 'Corrected_from': r.get('corrected_from'),
            'Data': r.get('Data', ''),
            **{m: r.get(m, np.nan) for m in MEASURES}, 'touse': bool(r.get('touse'))}


def write_analysis_xlsx(path, rows_by_subject, info=None, units=''):
    """Excel workbook: 'measures' (one row per trial) + the 12 summary sheets (+ 'info')."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'measures'
    ws.append(MEASURE_COLUMNS)
    for rows in rows_by_subject.values():
        for r in rows:
            ws.append([_cell(r.get(c)) for c in MEASURE_COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    for i in range(1, len(MEASURE_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    for name, table in summary_tables(rows_by_subject).items():
        sh = wb.create_sheet(name)
        cols = ['subject'] + [k for k in dict.fromkeys(k for row in table for k in row) if k != 'subject']
        sh.append(cols)
        for row in table:
            sh.append([_cell(row.get(c, np.nan)) for c in cols])
        for c in sh[1]:
            c.font = Font(bold=True)
        sh.freeze_panes = 'B2'
    sh = wb.create_sheet('info')
    lines = [('Times', 'seconds (responsetime, movementtime, AT, DT)'),
             ('Distances', units or 'see the analysis settings'),
             ('Method', 'Same as the previous MATLAB system (calculateMeasures.m): 100 Hz, 10 Hz Butterworth, '
                        '5% / 30% of peak speed thresholds'),
             ('Summaries', 'Only rows with touse = 1, familiarization excluded; sd = sample SD, iqr = MATLAB iqr')]
    for k, v in (info or {}).items():
        lines.append((k, str(v)))
    for line in lines:
        sh.append(list(line))
    sh.column_dimensions['A'].width = 18
    sh.column_dimensions['B'].width = 100
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
